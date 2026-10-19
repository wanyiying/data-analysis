"""Bounded local-only input; retain text IDs and reject malformed rows."""
import csv
import hashlib
import zipfile
from datetime import date, datetime
from xml.etree.ElementTree import ParseError
from pathlib import Path
from .contract import AuditError

MAX_BYTES = 20_000_000
MAX_ROWS = 100_000
MAX_COLUMNS = 100


def load(path, options=None):
    options = options or {}
    path = Path(path)
    if path.is_symlink() or any(p.is_symlink() for p in path.parents):
        raise AuditError("symlink_input", "ERROR")
    if not path.is_file() or path.suffix.lower() not in (".csv", ".xlsx"):
        raise AuditError("missing_or_unsupported_file", "ERROR")
    if path.stat().st_size > MAX_BYTES:
        raise AuditError("file_size_limit", "ERROR")
    digest = hashlib.sha256(path.read_bytes()).hexdigest()
    try:
        if path.suffix.lower() == ".csv":
            with path.open(encoding=options.get("encoding", "utf-8-sig"), newline="") as f:
                rows = list_bounded(csv.reader(f, strict=True))
        else:
            import openpyxl
            with zipfile.ZipFile(path) as z:
                if sum(i.file_size for i in z.infolist()) > MAX_BYTES * 5:
                    raise AuditError("excel_expansion_limit", "ERROR")
            wb = openpyxl.load_workbook(path, read_only=True, data_only=False)
            try:
                sheet = options.get("sheet")
                if sheet is None and len(wb.sheetnames) != 1:
                    raise AuditError("sheet_required", "ERROR")
                ws = wb[sheet or wb.sheetnames[0]]
                if ws.max_column and ws.max_column > MAX_COLUMNS:
                    raise AuditError("column_limit", "ERROR")
                rows = list_bounded(ws.values)
            finally:
                wb.close()
    except AuditError:
        raise
    except (OSError, UnicodeError, csv.Error, ValueError, LookupError, ParseError, zipfile.BadZipFile):
        raise AuditError("input_parse_error", "ERROR") from None
    if len(rows) < 2:
        raise AuditError("empty_data", "ERROR")
    header = rows[0]
    if any(not isinstance(h, str) or not h.strip() for h in header) or len(set(header)) != len(header):
        raise AuditError("invalid_header", "ERROR")
    if any(len(row) != len(header) for row in rows[1:]):
        raise AuditError("ragged_row", "ERROR")
    records = [dict(zip(header, ("" if v is None else v if isinstance(v, (date, datetime)) else str(v) for v in row))) for row in rows[1:]]
    return records, {"path": str(path.resolve()), "sha256": digest, "bytes": path.stat().st_size, "rows": len(records)}


def list_bounded(iterator):
    rows = []
    for row in iterator:
        if len(rows) > MAX_ROWS:
            raise AuditError("row_limit", "ERROR")
        if len(row) > MAX_COLUMNS:
            raise AuditError("column_limit", "ERROR")
        if any(len(str(v)) > 10000 for v in row if v is not None):
            raise AuditError("cell_size_limit", "ERROR")
        rows.append(list(row))
    return rows
