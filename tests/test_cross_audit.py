import copy
import csv
import json
import random
import subprocess
import sys
from decimal import Decimal
from pathlib import Path
import pytest

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "scripts"))
from reconcile.contract import AuditError
from reconcile.engine import audit
from reconcile.loader import load
from generate_cross_audit_cases import cases, base_contract


def records(rows):
    return [dict(zip(["id", "date", "region", "amount"], row)) for row in rows]


@pytest.mark.parametrize("name,a,b,c,expected", cases(), ids=[c[0] for c in cases()])
def test_known_answers(name, a, b, c, expected):
    if "error" in expected:
        with pytest.raises(AuditError) as exc:
            audit(records(a), records(b), c)
        assert exc.value.code == expected["error"]
        assert exc.value.state == expected.get("state", "BLOCKED")
        return
    result = audit(records(a), records(b), c)
    assert result["alignment"]["state"] == expected["state"]
    assert result["consistency"] == expected["consistency"]
    s = result["metrics"]["amount"]
    for expected_key, actual_key in (("net", "net_difference"), ("absolute", "absolute_difference"), ("only_a", "only_a_total"), ("only_b", "only_b_total")):
        if expected_key in expected:
            assert s[actual_key] == Decimal(expected[expected_key])
    assert s["identity_valid"]
    assert s["a_total"] == sum(Decimal(r[3]) * Decimal(c["sources"]["a"]["metrics"]["amount"]["multiplier"]) for r in a)
    assert s["b_total"] == sum(Decimal(r[3]) * Decimal(c["sources"]["b"]["metrics"]["amount"]["multiplier"]) for r in b)
    assert len(result["details"]) <= len(a)+len(b)
    assert result["correct_source"] == "not_determined"


def test_properties():
    rng = random.Random(7)
    for _ in range(40):
        a = [{"id": str(i), "date": "2026-01-01", "region": "east", "amount": str(rng.randint(-500, 500))} for i in range(10)]
        b = copy.deepcopy(a[2:])
        b[0]["amount"] = str(rng.randint(-500, 500))
        x = audit(a, b, base_contract())["metrics"]["amount"]
        y = audit(b, a, base_contract())["metrics"]["amount"]
        assert x["net_difference"] == -y["net_difference"]
        assert x["absolute_difference"] == y["absolute_difference"]
        rng.shuffle(a)
        assert audit(a, b, base_contract())["metrics"]["amount"] == x
        assert x["a_total"] == sum(Decimal(r["amount"]) for r in a)
        assert x["identity_valid"] and y["identity_valid"]


def test_split_invariance():
    c = base_contract()
    c.update(mode="aggregate", keys=["date", "region"])
    for s in c["sources"].values():
        s["grain"] = ["id", "date", "region"]
    a = records([["1", "2026-01-01", "east", "10"]])
    b = records([["2", "2026-01-01", "east", "4"], ["3", "2026-01-01", "east", "6"]])
    assert audit(a, b, c)["consistency"] == "consistent"


@pytest.mark.parametrize("content,code", [("", "empty_data"), ("id,amount\n", "empty_data"), ("id,id\na,1\n", "invalid_header"), ("id,amount\na,1,2\n", "ragged_row")])
def test_loader_errors(tmp_path, content, code):
    p = tmp_path / "x.csv"
    p.write_text(content, encoding="utf-8")
    with pytest.raises(AuditError, match=code):
        load(p)


def test_excel_ids_and_sheet(tmp_path):
    from openpyxl import Workbook
    wb = Workbook()
    wb.active.append(["id", "amount"])
    wb.active.append(["001", "1.01"])
    p = tmp_path / "x.xlsx"
    wb.save(p)
    assert load(p)[0][0]["id"] == "001"
    wb.create_sheet("other")
    wb.save(p)
    with pytest.raises(AuditError, match="sheet_required"):
        load(p)
    assert load(p, {"sheet": "Sheet"})[0][0]["amount"] == "1.01"


@pytest.mark.parametrize("value", ["NaN", "Infinity", "1e1000", "=SUM(A1)", "1,000", "garbage"])
def test_bad_amount(value):
    rows = records([["1", "2026-01-01", "east", value]])
    with pytest.raises(AuditError):
        audit(rows, rows, base_contract())


@pytest.mark.parametrize("mutation", [lambda c: c.update(row_tolerance="-1"), lambda c: c.update(mode="fuzzy"), lambda c: c["sources"]["a"]["fields"].pop("amount"), lambda c: c["sources"]["a"]["metrics"]["amount"].update(aggregation="distinct"), lambda c: c["sources"]["a"]["scope"].update(cutoff="yesterday")])
def test_invalid_config(mutation):
    c = base_contract()
    mutation(c)
    with pytest.raises(AuditError):
        audit([], [], c)


def test_multiple_metrics():
    c = base_contract()
    c["metrics"].append("units")
    for side in ("a", "b"):
        c["sources"][side]["fields"]["units"] = "units"
        c["sources"][side]["metrics"]["units"] = {"definition": "units", "unit": "item", "currency": "NA", "aggregation": "sum"}
    a = records([["1", "2026-01-01", "east", "10"]])
    a[0]["units"] = "2"
    b = copy.deepcopy(a)
    b[0]["units"] = "3"
    r = audit(a, b, c)
    assert r["metrics"]["amount"]["net_difference"] == 0
    assert r["metrics"]["units"]["net_difference"] == -1


@pytest.mark.parametrize("name,code", [("consistent", 0), ("value_mismatch", 1), ("definition", 3), ("confirmation", 4)])
def test_cli_e2e(tmp_path, name, code):
    case = ROOT / "examples" / "cross_audit" / name
    completed = subprocess.run([sys.executable, str(ROOT / "scripts/da_cross_reconcile.py"), "--source-a", str(case / "a.csv"), "--source-b", str(case / "b.csv"), "--config", str(case / "contract.json"), "--output-root", str(tmp_path / "runs")], capture_output=True, text=True)
    assert completed.returncode == code, completed.stderr
    run = Path(json.loads(completed.stdout)["run"])
    for filename in ("manifest.json", "profile.json", "contract.json", "alignment.json", "details.csv", "summary.json", "checks.json", "errors.json", "evidence.json", "agent-review.json", "verification.json", "report.html"):
        assert (run / filename).is_file()
    assert "DEGRADED" in (run / "report.html").read_text(encoding="utf-8")
    html = (run / "report.html").read_text(encoding="utf-8")
    if name == "value_mismatch":
        assert "net_difference" in html and "absolute_difference" in html
        from verify_cross_audit_run import verify
        assert verify(run)["status"] == "pass"
        original = json.loads((run / "summary.json").read_text(encoding="utf-8"))
        original["metrics"]["amount"]["net_difference"] = "99"
        (run / "summary.json").write_text(json.dumps(original), encoding="utf-8")
        assert verify(run)["status"] == "fail"


def test_input_hash_and_resource_limit(tmp_path, monkeypatch):
    import hashlib
    import reconcile.loader as loader
    p = tmp_path / "a.csv"
    p.write_text("id,amount\n001,2\n", encoding="utf-8")
    assert load(p)[1]["sha256"] == hashlib.sha256(p.read_bytes()).hexdigest()
    monkeypatch.setattr(loader, "MAX_BYTES", 1)
    with pytest.raises(AuditError, match="file_size_limit"):
        load(p)


def test_formula_export():
    from reconcile.evidence import safe_cell
    assert safe_cell("  =cmd()") == "'  =cmd()"
    assert safe_cell("-1") == "'-1"


def test_filter_and_date_scope():
    c = base_contract()
    for s in c["sources"].values():
        s["scope"]["filters"] = {"region": ["east"]}
    a = records([["1", "2026-01-01", "east", "2"], ["2", "2026-01-01", "west", "100"]])
    r = audit(a, a, c)
    assert r["metrics"]["amount"]["a_total"] == 2
    assert r["coverage"]["a"]["filtered"] == 1
    a[0]["date"] = "2026-02-01"
    with pytest.raises(AuditError, match="observed_date_outside_scope"):
        audit(a, a, c)


def test_evidence_adapter_numeric_binding(tmp_path):
    from reconcile.evidence import emit
    from da_verify import load_evidence
    c = base_contract()
    r = audit(records([["1", "2026-01-01", "east", "10"]]), records([["1", "2026-01-01", "east", "7"]]), c)
    emit(tmp_path, r, c, {}, {})
    evidence = json.loads((tmp_path / "evidence.json").read_text(encoding="utf-8"))
    assert evidence["evidence"][1]["exact_values"]["net_difference"] == "3"
    assert any(n[0] == 3 for n in load_evidence(evidence)["entries"]["X001"]["numbers"])


def test_html_escapes_untrusted_metric(tmp_path):
    from reconcile.evidence import emit
    c = base_contract()
    name = "<script>alert('injected')</script>"
    c["metrics"] = [name]
    for s in c["sources"].values():
        s["fields"][name] = s["fields"].pop("amount")
        s["metrics"][name] = s["metrics"].pop("amount")
    a = records([["1", "2026-01-01", "east", "10"]])
    emit(tmp_path, audit(a, a, c), c, {}, {})
    html = (tmp_path / "report.html").read_text(encoding="utf-8")
    assert "<script>alert('injected')</script>" not in html
    assert "&lt;script&gt;" in html


def test_public_sample_known_injection():
    case = ROOT / "examples/cross_audit/public_retail"
    if not case.exists():
        pytest.skip("Public sample has not been prepared")
    c = json.loads((case / "contract.json").read_text(encoding="utf-8"))
    r = audit(load(case / "a.csv")[0], load(case / "b.csv")[0], c)
    assert r["metrics"]["amount"]["net_difference"] == Decimal("-7.25")
    assert r["metrics"]["amount"]["absolute_difference"] == Decimal("7.25")
    assert r["metrics"]["amount"]["mismatch_count"] == 1


@pytest.mark.parametrize("reserved", ["__source_row__", "source_rows", "__metadata", "=amount", "+amount", "-amount", "@amount"])
def test_reserved_metric_prevents_false_consistency(reserved):
    c = base_contract()
    c["metrics"] = [reserved]
    for s in c["sources"].values():
        s["fields"][reserved] = s["fields"].pop("amount")
        s["metrics"][reserved] = s["metrics"].pop("amount")
    a = records([["1", "2026-01-01", "east", "10"]])
    b = records([["1", "2026-01-01", "east", "7"]])
    with pytest.raises(AuditError, match="reserved_or_invalid_field"):
        audit(a, b, c)


def test_native_excel_dates_equal_csv(tmp_path):
    from datetime import datetime
    from openpyxl import Workbook
    wb = Workbook()
    wb.active.append(["id", "date", "region", "amount"])
    wb.active.append(["001", datetime(2026, 1, 1), "east", "10.01"])
    p = tmp_path / "dates.xlsx"
    wb.save(p)
    csv_rows = records([["001", "2026-01-01", "east", "10.01"]])
    assert audit(load(p)[0], csv_rows, base_contract())["consistency"] == "consistent"


def test_no_output_overwrite_or_input_mutation(tmp_path):
    from reconcile.evidence import new_run
    root = tmp_path / "output"
    root.mkdir()
    prior = root / "history.txt"
    prior.write_text("keep", encoding="utf-8")
    first, second = new_run(root), new_run(root)
    assert first != second and first.is_dir() and second.is_dir()
    assert prior.read_text(encoding="utf-8") == "keep"
    a = records([["1", "2026-01-01", "east", "10"]])
    original = copy.deepcopy(a)
    audit(a, a, base_contract())
    assert a == original


def test_symlink_guard(tmp_path, monkeypatch):
    from reconcile.evidence import new_run
    monkeypatch.setattr(Path, "is_symlink", lambda self: self.name == "unsafe")
    with pytest.raises(AuditError, match="symlink_input"):
        load(tmp_path / "unsafe")
    with pytest.raises(ValueError, match="symlink_output"):
        new_run(tmp_path / "unsafe")


def test_row_and_column_limits(tmp_path, monkeypatch):
    import reconcile.loader as loader
    p = tmp_path / "data.csv"
    p.write_text("id,amount\na,1\nb,2\n", encoding="utf-8")
    monkeypatch.setattr(loader, "MAX_ROWS", 1)
    with pytest.raises(AuditError, match="row_limit"):
        load(p)
    monkeypatch.setattr(loader, "MAX_ROWS", 100)
    monkeypatch.setattr(loader, "MAX_COLUMNS", 1)
    with pytest.raises(AuditError, match="column_limit"):
        load(p)


def test_fixture_generation_reproducible(tmp_path, monkeypatch):
    import generate_cross_audit_cases as generator
    import hashlib
    monkeypatch.setattr(generator, "ROOT", tmp_path / "fixtures")
    generator.main()
    snapshot = {str(p.relative_to(generator.ROOT)): hashlib.sha256(p.read_bytes()).hexdigest() for p in generator.ROOT.rglob("*") if p.is_file()}
    generator.main()
    assert snapshot == {str(p.relative_to(generator.ROOT)): hashlib.sha256(p.read_bytes()).hexdigest() for p in generator.ROOT.rglob("*") if p.is_file()}


def test_gb18030_configured_csv(tmp_path):
    p = tmp_path / "encoded.csv"
    p.write_text("编号,金额\n001,2\n", encoding="gb18030")
    assert load(p, {"encoding": "gb18030"})[0][0]["编号"] == "001"
    with pytest.raises(AuditError, match="input_parse_error"):
        load(p, {"encoding": "does-not-exist"})


@pytest.mark.parametrize("bad_config", ["{bad", "[]", '{"schema_version":1,"sources":{"a":[]}}'])
def test_bad_config_still_produces_error_html(tmp_path, bad_config):
    config = tmp_path / "config.json"
    config.write_text(bad_config, encoding="utf-8")
    case = ROOT / "examples/cross_audit/consistent"
    completed = subprocess.run([sys.executable, str(ROOT / "scripts/da_cross_reconcile.py"), "--source-a", str(case / "a.csv"), "--source-b", str(case / "b.csv"), "--config", str(config), "--output-root", str(tmp_path / "runs")], capture_output=True, text=True)
    assert completed.returncode == 2
    result = json.loads(completed.stdout)
    assert result["execution"] == "error"
    assert (Path(result["run"]) / "report.html").is_file()
